# -*- coding: utf-8 -*-
import openpyxl
import csv
import re
import glob
import os
from docx.shared import Cm 
from docx import Document
import docx 
from openpyxl import Workbook

Source_files = glob.glob(os.getcwd() + "/tN_Marathi/*.docx")
print(Source_files)
# Source_files = glob.glob(os.getcwd() + "/Source/*.tsv")
bookdict = { "2CO": "en_tn_48-2CO", "JUD": "en_tn_66-JUD", "2PE": "en_tn_62-2PE", "GAL": "en_tn_49-GAL", "JHN": "en_tn_44-JHN", "PHP": "en_tn_51-PHP", "LUK": "en_tn_43-LUK", "2TI": "en_tn_56-2TI", "MAT": "en_tn_41-MAT", "ACT": "en_tn_45-ACT", "PHM": "en_tn_58-PHM", "HEB": "en_tn_59-HEB", "JAS": "en_tn_60-JAS", "TIT": "en_tn_57-TIT", "COL": "en_tn_52-COL", "ROM": "en_tn_46-ROM", "1JN": "en_tn_63-1JN", "1TH": "en_tn_53-1TH", "1TI": "en_tn_55-1TI", "MRK": "en_tn_42-MRK", "2JN": "en_tn_64-2JN", "1PE": "en_tn_61-1PE", "2TH": "en_tn_54-2TH", "REV": "en_tn_67-REV", "3JN": "en_tn_65-3JN", "EPH": "en_tn_50-EPH", "1CO": "en_tn_47-1CO"}
# targetPath = glob.glob(os.getcwd() + "/Source/*.tsv") 
for s_file in Source_files:
    bookName = s_file.split("/")[-1].split(".")[0].split("-")[-1].split()[0]
    print(bookName)
    document = Document(s_file)
    tables = document.tables
    xlxbook = Workbook()
    sheet = xlxbook.active 
    if bookName in bookdict:
        fileName  = bookdict.get(bookName)
    else:
        fileName = bookName

    for table in tables:
        tl_rows = table.rows
        for row in tl_rows:
            try:
                tl_book = row.cells[0].text
                tl_chapter = row.cells[1].text
                tl_verse = row.cells[2].text
                tl_english = row.cells[3].text
                tl_occur = row.cells[4].text
                sheet.append((tl_book,tl_chapter,tl_verse,tl_english,tl_occur))
            except:
                pass
    
    xlxbook.save(fileName+'.xlsx')
            


